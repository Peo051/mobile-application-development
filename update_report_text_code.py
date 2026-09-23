import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc_path = r"f:\Thuc Hanh\LTDD\docs\BaiTapTuan5_TranDuongGiaBao_2001240039.docx"
screenshots_dir = r"f:\Thuc Hanh\LTDD\screenshots"

doc = Document()

# Thiết lập lề trang chuẩn A4
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor(26, 82, 118)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(90, 90, 90)

def add_info_box():
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    shading = parse_xml(r'<w:shd {} w:fill="F4F6F7"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="8" w:space="0" w:color="BDC3C7"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1A5276"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="BDC3C7"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="BDC3C7"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    lines = [
        ("Họ và tên: ", "Trần Dương Gia Bảo"),
        ("Mã số sinh viên: ", "2001240039"),
        ("Lớp học phần: ", "15DHTH02"),
        ("Môn học: ", "Lập trình di động (Flutter & Dart)"),
        ("Github: ", "https://github.com/Peo051/mobile-application-development"),
    ]
    for label, val in lines:
        r1 = p.add_run(label)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        
        r2 = p.add_run(val + "\n" if label != lines[-1][0] else val)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        if "http" in val:
            r2.font.color.rgb = RGBColor(21, 101, 192)
            r2.font.underline = True

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(180, 40, 40)

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(24, 106, 59)

def add_desc(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(80, 80, 80)

def add_sublabel(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(40, 40, 40)

# Hàm chèn Code dạng Text trong khung bảng có nền và viền đẹp mắt
def add_code_block(code_text, file_tag=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    # Nền xám nhạt như GitHub code block (#F6F8FA)
    shading = parse_xml(r'<w:shd {} w:fill="F6F8FA"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    # Border viền trái xanh đậm, các cạnh khác xám nhạt
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="0969DA"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    
    # Đệm lề bên trong
    tcMar = parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="180" w:type="dxa"/>
            <w:right w:w="180" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if file_tag:
        r_tag = p.add_run(f"// --- {file_tag} ---\n")
        r_tag.font.name = "Consolas"
        r_tag.font.size = Pt(8.5)
        r_tag.font.bold = True
        r_tag.font.color.rgb = RGBColor(9, 105, 218)
    
    r_code = p.add_run(code_text.strip())
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(36, 41, 47)

def add_image_centered(img_name, width_inch=2.8):
    path = os.path.join(screenshots_dir, img_name)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.add_run().add_picture(path, width=Inches(width_inch))
    else:
        print(f"Warning: Image not found: {path}")

# ==================== NỘI DUNG TÀI LIỆU ====================

add_title("BÁO CÁO BÀI TẬP THỰC HÀNH - TUẦN 5")
add_subtitle("CHƯƠNG 03: THIẾT KẾ LAYOUT TRONG FLUTTER")
add_info_box()

# ----------------- BÀI TẬP 03 -----------------
add_heading_1("BÀI TẬP 03: THIẾT KẾ GIAO DIỆN \"QUÀ CỦA VINH (7)\"")
add_desc("Mô phỏng màn hình Quà của Vinh trên ứng dụng Momo: Thanh tiêu đề màu hồng nhạt, bộ lọc cuộn ngang, 2 thẻ thông tin điểm/ưu đãi, và danh sách 5 quà tặng thương hiệu thực tế (CGV, Sim Mobifone, Ngân hàng VIB, Bảo hiểm, Phí không dừng).")

add_sublabel("1. Source code tạo nên giao diện Quà của Vinh:")
code_bai03 = """  @override
  Widget build(BuildContext context) {
    return Scaffold(
      backgroundColor: screenBackground,
      body: SafeArea(
        child: Column(
          children: [
            _buildAppBar(), // Thanh tiêu đề "Quà của Vinh (7)" + nút back + hotline
            Expanded(
              child: SingleChildScrollView(
                child: Column(
                  crossAxisAlignment: CrossAxisAlignment.start,
                  children: [
                    const SizedBox(height: 10),
                    _buildFilterBar(),     // Thanh bộ lọc cuộn ngang (Sắp xếp, Dịch vụ, ...)
                    const SizedBox(height: 12),
                    _buildTwoInfoCards(),  // 2 thẻ thông tin nổi bật (Điểm tích luỹ & Ưu đãi)
                    const SizedBox(height: 14),
                    _buildGiftList(),      // Danh sách 5 mục quà tặng với logo CGV, VIB, ...
                    const SizedBox(height: 20),
                  ],
                ),
              ),
            ),
          ],
        ),
      ),
    );
  }

  /// Widget hiển thị từng thẻ Quà tặng với logo, tiêu đề, tag và nút thao tác
  Widget _buildGiftCard(GiftItem item) {
    return Container(
      margin: const EdgeInsets.symmetric(horizontal: 14, vertical: 6),
      padding: const EdgeInsets.all(12),
      decoration: BoxDecoration(
        color: Colors.white,
        borderRadius: BorderRadius.circular(14),
        boxShadow: [
          BoxShadow(
            color: Colors.black.withOpacity(0.04),
            blurRadius: 6,
            offset: const Offset(0, 2),
          ),
        ],
      ),
      child: Row(
        children: [
          // Logo thương hiệu (CGV, VIB, Bảo hiểm, Sim, Phí không dừng)
          ClipRRect(
            borderRadius: BorderRadius.circular(10),
            child: Image.asset(item.assetPath, width: 56, height: 56, fit: BoxFit.contain),
          ),
          const SizedBox(width: 12),
          // Thông tin khuyến mãi, tiêu đề, HSD
          Expanded(
            child: Column(
              crossAxisAlignment: CrossAxisAlignment.start,
              children: [
                Text(item.brandName, style: const TextStyle(fontWeight: FontWeight.bold, fontSize: 13)),
                Text(item.title, style: const TextStyle(fontWeight: FontWeight.w600, fontSize: 15)),
                Text(item.description, style: TextStyle(color: Colors.grey.shade600, fontSize: 12)),
                if (item.expiryDate != null)
                  Text(item.expiryDate!, style: TextStyle(color: Colors.grey.shade500, fontSize: 11)),
              ],
            ),
          ),
          // Nút Thu thập hoặc biểu tượng Trái tim yêu thích
          if (item.isCollectButton)
            OutlinedButton(
              onPressed: () {},
              style: OutlinedButton.styleFrom(foregroundColor: primaryPink),
              child: const Text('Thu thập'),
            )
          else
            Icon(item.isFavorite ? Icons.favorite : Icons.favorite_border,
                 color: item.isFavorite ? primaryPink : Colors.grey),
        ],
      ),
    );
  }"""
add_code_block(code_bai03, file_tag="lib/bai_03.dart")

add_sublabel("2. Ảnh chụp kết quả giao diện thực tế:")
add_image_centered("week5_bai3_qua_cua_vinh.png", width_inch=3.0)


# ----------------- BÀI TẬP 04 -----------------
doc.add_page_break()
add_heading_1("BÀI TẬP 04: THIẾT KẾ ỨNG DỤNG CỬA HÀNG ĐIỆN THOẠI ĐƠN GIẢN")
add_desc("Xây dựng ứng dụng hoàn chỉnh gồm: Màn hình giới thiệu (Intro), Màn hình Cửa hàng (Shop) danh sách điện thoại cuộn ngang, Menu Drawer điều hướng, Hộp thoại thêm vào giỏ, Màn hình Giỏ hàng (Cart) với chức năng xóa và hộp thoại thanh toán.")

# Mục 4.1
add_heading_2("1. Màn hình Giới thiệu cửa hàng (IntroScreen)")
add_desc("Màn hình mở đầu với logo HUIT tròn, tên cửa hàng, địa chỉ và nút chuyển tiếp bo tròn viên thuốc.")
add_sublabel("a) Source code tạo màn hình Giới thiệu:")
code_intro = """class IntroScreen extends StatelessWidget {
  const IntroScreen({super.key});

  @override
  Widget build(BuildContext context) {
    return Scaffold(
      backgroundColor: const Color(0xFFE0E0E0),
      body: SafeArea(
        child: Center(
          child: Padding(
            padding: const EdgeInsets.symmetric(horizontal: 32),
            child: Column(
              mainAxisAlignment: MainAxisAlignment.center,
              children: [
                // Logo trường HUIT tròn lớn ở giữa màn hình
                Container(
                  width: 170,
                  height: 170,
                  decoration: const BoxDecoration(shape: BoxShape.circle),
                  child: ClipOval(
                    child: Image.asset('assets/images/logo_huit.png', fit: BoxFit.contain),
                  ),
                ),
                const SizedBox(height: 32),
                // Tiêu đề Cửa hàng điện thoại
                const Text(
                  'Cửa hàng điện thoại',
                  style: TextStyle(fontSize: 22, fontWeight: FontWeight.bold),
                  textAlign: TextAlign.center,
                ),
                const SizedBox(height: 12),
                // Địa chỉ cửa hàng
                const Text(
                  '140 Lê Trọng Tấn, Tân Phú, TP.Hồ Chí Minh',
                  style: TextStyle(fontSize: 13.5, color: Colors.black54),
                  textAlign: TextAlign.center,
                ),
                const SizedBox(height: 36),
                // Nút hình viên thuốc màu trắng chuyển sang ShopScreen
                InkWell(
                  onTap: () {
                    Navigator.push(
                      context,
                      MaterialPageRoute(builder: (context) => const ShopScreen()),
                    );
                  },
                  borderRadius: BorderRadius.circular(30),
                  child: Container(
                    padding: const EdgeInsets.symmetric(horizontal: 26, vertical: 12),
                    decoration: BoxDecoration(
                      color: Colors.white,
                      borderRadius: BorderRadius.circular(30),
                      boxShadow: [
                        BoxShadow(
                          color: Colors.black.withOpacity(0.08),
                          blurRadius: 10,
                          offset: const Offset(0, 4),
                        ),
                      ],
                    ),
                    child: const Icon(Icons.arrow_forward, color: Colors.black54, size: 22),
                  ),
                ),
              ],
            ),
          ),
        ),
      ),
    );
  }
}"""
add_code_block(code_intro, file_tag="lib/screens/intro_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_01_intro.png", width_inch=2.8)

# Mục 4.2
add_heading_2("2. Màn hình Cửa hàng (ShopScreen - Danh sách sản phẩm)")
add_desc("Giao diện Cửa hàng điện thoại với AppBar màu vàng cam, danh sách sản phẩm hiển thị dưới dạng thẻ cuộn ngang kèm nút thêm vào giỏ hàng.")
add_sublabel("a) Source code bố cục Cửa hàng & Danh sách sản phẩm:")
code_shop = """  @override
  Widget build(BuildContext context) {
    return Scaffold(
      backgroundColor: const Color(0xFFE0E0E0),
      appBar: AppBar(
        backgroundColor: const Color(0xFFFFC107), // Màu vàng cam theo thiết kế
        leading: Builder(
          builder: (context) => IconButton(
            icon: const Icon(Icons.menu, color: Colors.black87),
            onPressed: () => Scaffold.of(context).openDrawer(),
          ),
        ),
        title: const Text('Cửa hàng điện thoại',
            style: TextStyle(color: Colors.black87, fontWeight: FontWeight.bold)),
        actions: [
          IconButton(
            icon: const Icon(Icons.shopping_cart, color: Colors.black87),
            onPressed: _openCartScreen,
          ),
        ],
      ),
      drawer: Drawer(...), // Ngăn kéo menu bên trái
      body: SafeArea(
        child: SingleChildScrollView(
          child: Column(
            children: [
              const SizedBox(height: 16),
              const Text('Chọn sản phẩm bạn muốn sử dụng',
                  style: TextStyle(fontSize: 14.5, fontWeight: FontWeight.bold)),
              const SizedBox(height: 16),
              // ListView cuộn ngang chứa danh sách sản phẩm
              SizedBox(
                height: 380,
                child: ListView.builder(
                  scrollDirection: Axis.horizontal,
                  padding: const EdgeInsets.only(left: 18),
                  itemCount: sampleProducts.length,
                  itemBuilder: (context, index) {
                    final product = sampleProducts[index];
                    return ProductCard(
                      product: product,
                      onAddToCart: () => _showAddToCartDialog(product),
                    );
                  },
                ),
              ),
              const SizedBox(height: 20),
              const Text('Sản phẩm được lựa chọn nhiều nhất',
                  style: TextStyle(fontSize: 13.5, fontWeight: FontWeight.bold)),
            ],
          ),
        ),
      ),
    );
  }"""
add_code_block(code_shop, file_tag="lib/screens/shop_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_02_shop.png", width_inch=2.8)

# Mục 4.3
add_heading_2("3. Ngăn kéo Menu điều hướng (Drawer Navigation)")
add_desc("Menu bên trái gồm Header màu trắng chứa ảnh đại diện, tên giảng viên Vũ Văn Vĩnh, email, và thân menu màu xanh dương với các mục Cửa hàng, Giỏ hàng, Thoát.")
add_sublabel("a) Source code cấu hình Drawer:")
code_drawer = """      drawer: Drawer(
        child: Column(
          children: [
            // Phần trên: Thông tin cá nhân nền trắng
            Container(
              color: Colors.white,
              width: double.infinity,
              padding: const EdgeInsets.fromLTRB(20, 48, 20, 20),
              child: Column(
                children: [
                  ClipOval(
                    child: Image.asset('assets/images/logo_huit.png',
                        width: 72, height: 72, fit: BoxFit.contain),
                  ),
                  const SizedBox(height: 12),
                  const Text('Vũ Văn Vĩnh',
                      style: TextStyle(fontSize: 16, fontWeight: FontWeight.bold)),
                  const SizedBox(height: 4),
                  Text('vinhvv@huit.edu.vn',
                      style: TextStyle(fontSize: 13, color: Colors.grey.shade600)),
                ],
              ),
            ),
            // Phần dưới: Menu điều hướng nền màu xanh dương
            Expanded(
              child: Material(
                color: const Color(0xFF2196F3),
                child: Padding(
                  padding: const EdgeInsets.symmetric(vertical: 8),
                  child: Column(
                    children: [
                      ListTile(
                        leading: const Icon(Icons.store, color: Colors.black87),
                        title: const Text('Cửa hàng'),
                        onTap: () => Navigator.pop(context),
                      ),
                      ListTile(
                        leading: const Icon(Icons.shopping_cart, color: Colors.black87),
                        title: const Text('Giỏ hàng'),
                        onTap: () {
                          Navigator.pop(context);
                          _openCartScreen();
                        },
                      ),
                      ListTile(
                        leading: const Icon(Icons.exit_to_app, color: Colors.black87),
                        title: const Text('Thoát'),
                        onTap: () {
                          Navigator.pop(context);
                          Navigator.popUntil(context, (route) => route.isFirst);
                        },
                      ),
                    ],
                  ),
                ),
              ),
            ),
          ],
        ),
      ),"""
add_code_block(code_drawer, file_tag="lib/screens/shop_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_03_drawer.png", width_inch=2.8)

# Mục 4.4
add_heading_2("4. Hộp thoại Xác nhận thêm sản phẩm vào Giỏ hàng")
add_desc("Hộp thoại AlertDialog xuất hiện khi người dùng nhấn nút dấu cộng (+) màu xanh ngọc trên từng thẻ sản phẩm.")
add_sublabel("a) Source code hộp thoại Xác nhận thêm vào giỏ:")
code_dialog_add = """  // Hiển thị hộp thoại xác nhận khi người dùng nhấn nút '+' trên thẻ sản phẩm
  void _showAddToCartDialog(Product product) {
    showDialog(
      context: context,
      builder: (ctx) => AlertDialog(
        shape: RoundedRectangleBorder(borderRadius: BorderRadius.circular(16)),
        title: const Text(
          'Xác nhận',
          style: TextStyle(fontWeight: FontWeight.bold),
        ),
        content: const Text('Bạn vừa thêm sản phẩm vào Giỏ hàng'),
        actions: [
          TextButton(
            onPressed: () => Navigator.pop(ctx),
            child: const Text('Không', style: TextStyle(color: Colors.black87)),
          ),
          TextButton(
            onPressed: () {
              Navigator.pop(ctx);
              setState(() {
                _cartItems.add(product); // Thêm sản phẩm vào danh sách giỏ hàng
              });
              ScaffoldMessenger.of(context).showSnackBar(
                SnackBar(
                  content: Text('Đã thêm ${product.name} vào Giỏ hàng'),
                  duration: const Duration(seconds: 1),
                  behavior: SnackBarBehavior.floating,
                ),
              );
            },
            child: const Text(
              'Đồng ý',
              style: TextStyle(color: Colors.black87, fontWeight: FontWeight.bold),
            ),
          ),
        ],
      ),
    );
  }"""
add_code_block(code_dialog_add, file_tag="lib/screens/shop_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_04_dialog_add.png", width_inch=2.8)

# Mục 4.5
add_heading_2("5. Màn hình Giỏ hàng của bạn (CartScreen)")
add_desc("Màn hình hiển thị danh sách các điện thoại đã chọn, mỗi mục có tên, giá và icon thùng rác để xóa, cùng nút 'Thanh toán' ở đáy màn hình.")
add_sublabel("a) Source code màn hình Giỏ hàng:")
code_cart = """  @override
  Widget build(BuildContext context) {
    final bool isEmpty = widget.cartItems.isEmpty;

    return Scaffold(
      backgroundColor: const Color(0xFFE0E0E0),
      appBar: AppBar(
        backgroundColor: const Color(0xFFFFC107),
        leading: IconButton(
          icon: const Icon(Icons.arrow_back, color: Colors.black87),
          onPressed: () => Navigator.pop(context),
        ),
        title: const Text('Giỏ hàng của bạn',
            style: TextStyle(color: Colors.black87, fontWeight: FontWeight.bold)),
      ),
      body: SafeArea(
        child: Column(
          children: [
            const SizedBox(height: 12),
            const Text('Giỏ hàng của bạn',
                style: TextStyle(fontSize: 12.5, color: Colors.black87)),
            const SizedBox(height: 12),
            // Danh sách các sản phẩm đã thêm kèm nút thùng rác xóa sản phẩm
            Expanded(
              child: isEmpty
                  ? const Center(
                      child: Text('Bạn chưa có sản phẩm nào vô giỏ hàng!!!!',
                          style: TextStyle(fontSize: 14.5, fontWeight: FontWeight.w500)),
                    )
                  : ListView.builder(
                      padding: const EdgeInsets.symmetric(horizontal: 24),
                      itemCount: widget.cartItems.length,
                      itemBuilder: (context, index) {
                        final item = widget.cartItems[index];
                        return Padding(
                          padding: const EdgeInsets.symmetric(vertical: 8),
                          child: Row(
                            mainAxisAlignment: MainAxisAlignment.spaceBetween,
                            children: [
                              Column(
                                crossAxisAlignment: CrossAxisAlignment.start,
                                children: [
                                  Text(item.name,
                                      style: const TextStyle(fontWeight: FontWeight.bold)),
                                  Text(item.price.toString()),
                                ],
                              ),
                              IconButton(
                                icon: const Icon(Icons.delete, color: Colors.black87, size: 20),
                                onPressed: () => _confirmRemove(index, item.name),
                              ),
                            ],
                          ),
                        );
                      },
                    ),
            ),
            // Nút Thanh toán hình viên thuốc bo tròn ở chân trang
            Padding(
              padding: const EdgeInsets.only(bottom: 24, top: 12),
              child: InkWell(
                onTap: _handleCheckout,
                borderRadius: BorderRadius.circular(30),
                child: Container(
                  padding: const EdgeInsets.symmetric(horizontal: 36, vertical: 10),
                  decoration: BoxDecoration(
                    color: Colors.white,
                    borderRadius: BorderRadius.circular(30),
                  ),
                  child: const Text('Thanh toán',
                      style: TextStyle(color: Color(0xFF00D1C7), fontWeight: FontWeight.bold)),
                ),
              ),
            ),
          ],
        ),
      ),
    );
  }"""
add_code_block(code_cart, file_tag="lib/screens/cart_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_05_cart_with_items.png", width_inch=2.8)

# Mục 4.6
add_heading_2("6. Hộp thoại Xác nhận loại bỏ sản phẩm ra khỏi giỏ hàng")
add_desc("Hộp thoại AlertDialog cảnh báo khi người dùng nhấn vào biểu tượng thùng rác của sản phẩm trong giỏ.")
add_sublabel("a) Source code hộp thoại Xác nhận loại bỏ sản phẩm:")
code_dialog_remove = """  // Hiển thị hộp thoại xác nhận khi bấm biểu tượng thùng rác để xóa sản phẩm
  void _confirmRemove(int index, String productName) {
    showDialog(
      context: context,
      builder: (ctx) => AlertDialog(
        shape: RoundedRectangleBorder(borderRadius: BorderRadius.circular(16)),
        title: const Text(
          'Xác nhận',
          style: TextStyle(fontWeight: FontWeight.bold),
        ),
        content: const Text('Bạn muốn loại bỏ sản phẩm này ra khỏi giỏ hàng'),
        actions: [
          TextButton(
            onPressed: () => Navigator.pop(ctx),
            child: const Text('Không', style: TextStyle(color: Colors.black87)),
          ),
          TextButton(
            onPressed: () {
              Navigator.pop(ctx);
              setState(() {
                widget.onRemoveItem(index); // Xóa sản phẩm tại vị trí index
              });
            },
            child: const Text(
              'Đồng ý',
              style: TextStyle(color: Colors.black87, fontWeight: FontWeight.bold),
            ),
          ),
        ],
      ),
    );
  }"""
add_code_block(code_dialog_remove, file_tag="lib/screens/cart_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_07_dialog_remove.png", width_inch=2.8)

# Mục 4.7
add_heading_2("7. Hộp thoại Xác nhận Thanh toán thành công")
add_desc("Hộp thoại AlertDialog thông báo 'Bạn đã thanh toán xong giỏ hàng' khi người dùng nhấn nút Thanh toán, sau đó làm rỗng giỏ hàng.")
add_sublabel("a) Source code hộp thoại Thanh toán:")
code_dialog_checkout = """  // Xử lý khi nhấn nút Thanh toán
  void _handleCheckout() {
    if (widget.cartItems.isEmpty) {
      ScaffoldMessenger.of(context).showSnackBar(
        const SnackBar(
          content: Text('Bạn chưa có sản phẩm nào vô giỏ hàng!!!!'),
          duration: Duration(seconds: 2),
        ),
      );
      return;
    }

    // Hiển thị dialog thông báo đã thanh toán xong và xóa sạch giỏ hàng
    showDialog(
      context: context,
      builder: (ctx) => AlertDialog(
        shape: RoundedRectangleBorder(borderRadius: BorderRadius.circular(16)),
        title: const Text(
          'Thanh toán',
          style: TextStyle(fontWeight: FontWeight.bold),
        ),
        content: const Text('Bạn đã thanh toán xong giỏ hàng'),
        actions: [
          TextButton(
            onPressed: () {
              Navigator.pop(ctx);
              setState(() {
                widget.onClearCart(); // Làm rỗng giỏ hàng sau khi thanh toán
              });
            },
            child: const Text(
              'Đồng ý',
              style: TextStyle(color: Colors.black87, fontWeight: FontWeight.bold),
            ),
          ),
        ],
      ),
    );
  }"""
add_code_block(code_dialog_checkout, file_tag="lib/screens/cart_screen.dart")

add_sublabel("b) Kết quả giao diện thực tế:")
add_image_centered("week5_bai4_06_checkout.png", width_inch=2.8)

doc.save(doc_path)
print(f"Report updated successfully with TEXT CODE at: {doc_path}")
