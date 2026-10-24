from docx import Document
import sys

path = sys.argv[1]
doc = Document(path)
cell = doc.tables[4].cell(0, 0)
assert cell.paragraphs[0].text == 'Lưu ý về schema'
assert 'chưa được scaffold hoặc đối chiếu' in cell.paragraphs[1].text
cell.paragraphs[0].runs[0].text = 'Schema đã xác nhận'
cell.paragraphs[1].text = ('Database QuanLyDoAn đã được tạo từ QuanLyDoAn.sql trên instance LAPTOP-VHGPK0SP. '
                            'Dữ liệu mẫu và khóa ngoại đã được kiểm tra; DBCC CHECKCONSTRAINTS không báo lỗi, '
                            'và các endpoint đọc dữ liệu hiện trả HTTP 200.')
doc.save(path)
print('Updated schema status in', path)
